"""Reapply the reviewed changes to the original reports without changing the template.

Run from the repository root with python-docx and Pillow installed.
Reads the retained Git baseline, the measured evidence and the role-specific edits.
Only the requested role is rewritten. Word rendering is a separate shipping gate.
"""
import argparse
import io
import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'evidence/2026-10-27'
BASELINE = '891daec'

ROLE_CONTENT = {
 'A': {
  'functions': '本模块接收本轮配置、名单ID和一个选项索引，输出固定维度的随机密文向量。创建流程生成2048位Paillier密钥，登记流程生成Ed25519签名密钥，口令保护的保管库保存私钥材料。选择为单选，选项数量可以变化，选票中不保存明文索引。',
  'design': '先用独热向量约束客户端表达方式，再逐分量调用raw_encrypt；把公开随机标识nonce与加密内部随机数r区分，避免公开r。选项维度来自配置，不把三选项写死。密钥保管采用scrypt与AES-GCM，错误口令通过认证标签校验失败返回。',
  'flow': ['输入配置、名单ID和选择索引', 'ID在名单且索引合法？', '构造独热向量，每项为0或1', '用独立随机数加密各分量', '封装轮次、票型和时间|附加公开随机标识nonce', '交给登记端校验并接收'],
  'issue': '新增边界实验表明，客户端正确生成独热向量不等于接收端能证明密文合法。两名恶意客户端分别提交[2,0]与[0,0]的加密向量，最终总票为[2,0]，仍满足总和等于两票。处理方法是保留总和检查作为有限检测，明确合法性证明尚未实现，不把它写成“所有非法票都能被发现”。',
  'extra': [('非法选票组合局限', '两个非独热向量仍可形成满足约束的汇总；需零知识证明')],
 },
 'B': {
  'functions': '计票模块只需公开配置、密文选票和签名清单，输出每个选项的密文乘积。保管端核验当前关闭状态及完整汇总后，使用私钥返回总票数。两端通过JSON请求关联轮次、清单摘要、密文向量与接收票数，结果须满足非负、单项上界和总和约束。',
  'design': '选择按选项列进行模乘，避免逐票解密。累积元取1，因此空集合解密为0。为避免解密端成为任意密文查询服务，在加载私钥前独立审计和复算；对不完整或过期请求明确拒绝。耗时由复算端自身测量，不能由调用方提供负数或伪造值。',
  'flow': ['读取公开配置、密文与清单', '审计通过且投票已关闭？', '按选项列累乘密文并对n²取模', '保管端独立复算并核对请求', '请求匹配且总票约束成立？', '仅返回各选项总票和本次耗时'],
  'issue': '审查时发现，解密端虽然重新计算了完整汇总，却把请求内的audit_ms和aggregate_ms直接写入结果。调用者可传入负值污染性能记录。修复后结果改取expected中的复算耗时；回归测试用-999作为请求耗时，最终票数仍为[1,0]且返回耗时均为非负。',
  'extra': [('单人投票隐私局限', '一人总票[0,1]直接暴露该人的选择'), ('非法选票组合局限', '总和约束不能替代逐票独热合法性证明')],
 },
 'C': {
  'functions': '界面提供创建、投票、结果和审计四个页面；登记服务负责名单、唯一ID、关闭状态及签名清单。审计模块校验JSON结构、签名、配置摘要和选票集合；导出功能只交付当前已审计的公开文件。后台线程执行耗时操作，主线程更新控件。',
  'design': '先固定公开JSON接口，再把每张选票摘要加入签名清单，使修改与删除都有可比较的基准。文件损坏时统一返回可解释的VoteError，并拒绝继续计票。公开包采用明确文件清单，避免递归搜集目录中的无关JSON；GUI通过单写入锁和任务禁用防止常规并发操作。',
  'flow': ['选择轮次或提交新的加密选票', '结构、签名与名单校验通过？', '检查重复ID、nonce与关闭状态', '写入选票并签署最新清单', '关闭后审计并调用完整计票', '显示汇总或导出已审计公开文件'],
  'issue': '公开导出原先递归收集public下的全部JSON。复现时在该目录加入accidental-secret.json，文件也进入ZIP。修复后仅导出config.json、manifest.json和签名清单列出的选票；额外顶层JSON与备份子目录不再外发。另将损坏JSON、布尔版本号和空对象改为明确拒绝，并使选票子目录中的未登记文件触发审计失败。',
  'extra': [('公开包白名单回归', '无关JSON不进入ZIP，无私钥目录仍可复算相同汇总'), ('损坏输入回归', '损坏JSON、空对象及布尔版本返回VoteError')],
 }
}


def font(run, size=12, bold=False, east='宋体', latin='Times New Roman'):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = __import__('docx.shared', fromlist=['RGBColor']).RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), east)


def replace(p, text):
    p.clear()
    p.add_run(text)


def insert_before(anchor, text, style='Report Body'):
    p = anchor.insert_paragraph_before(text, style)
    return p


def insert_after(anchor, text, style='Report Body'):
    el = OxmlElement('w:p')
    anchor._p.addnext(el)
    p = __import__('docx.text.paragraph', fromlist=['Paragraph']).Paragraph(el, anchor._parent)
    p.style = style
    p.add_run(text)
    return p


def styles(doc):
    for name, size, bold in [('Report Title',16,True), ('Report H1',14,True),
                              ('Report H2',12,True), ('Report Body',12,False),
                              ('Report Caption',10.5,False), ('Report Code',10,False)]:
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = doc.styles['Normal']
        s.font.name = 'Times New Roman'
        s.font.size = Pt(size)
        s.font.bold = bold
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
        fmt = s.paragraph_format
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)
        fmt.line_spacing = 1.5
        fmt.first_line_indent = Pt(24 if name == 'Report Body' else 0)
        fmt.widow_control = True
        if name == 'Report Body':
            fmt.keep_together = True
        if name in ('Report Title', 'Report H1', 'Report H2'):
            fmt.keep_with_next = True
            fmt.space_before = Pt(6)
            fmt.space_after = Pt(3)
        if name in ('Report H1', 'Report H2'):
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), '0' if name == 'Report H1' else '1')
            s.element.get_or_add_pPr().append(outline)
        if name == 'Report Caption':
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.space_after = Pt(6)
        if name == 'Report Code':
            fmt.line_spacing = 1
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.keep_together = True
            fmt.left_indent = Pt(12)
            s.font.name = 'Consolas'
            s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '新宋体')


def flowchart(role, output):
    # A program diagram authored from source logic, not a generated illustration.
    labels = ROLE_CONTENT[role]['flow']
    img = Image.new('RGB', (1500, 900), 'white')
    draw = ImageDraw.Draw(img)
    f = ImageFont.truetype('C:/Windows/Fonts/simsun.ttc', 36)
    small = ImageFont.truetype('C:/Windows/Fonts/simsun.ttc', 29)
    centers = [(430,105),(430,365),(430,625),(1060,625),(1060,365),(1060,105)]
    def arrow(a, b):
        draw.line([a,b],fill='black',width=4)
        x,y=b; dx=b[0]-a[0];dy=b[1]-a[1]
        if abs(dx)>abs(dy):
            s=1 if dx>0 else -1;draw.polygon([(x,y),(x-15*s,y-9),(x-15*s,y+9)],fill='black')
        else:
            s=1 if dy>0 else -1;draw.polygon([(x,y),(x-9,y-15*s),(x+9,y-15*s)],fill='black')
    for i,((x,y),label) in enumerate(zip(centers,labels)):
        if i == 1:
            draw.polygon([(x,y-95),(x+275,y),(x,y+95),(x-275,y)],outline='black',width=3)
        else:
            draw.rounded_rectangle((x-285,y-65,x+285,y+65),radius=13,outline='black',width=3)
        parts = label.split('|') if '|' in label else ([label] if len(label)<=15 else [label[:15],label[15:]])
        for j,line in enumerate(parts):
            box=draw.textbbox((0,0),line,font=f)
            draw.text((x-(box[2]-box[0])/2,y-24+(j-(len(parts)-1)/2)*42),line,font=f,fill='black')
    arrow((430,170),(430,270)); arrow((430,460),(430,560)); arrow((715,625),(775,625))
    arrow((1060,560),(1060,430)); arrow((1060,300),(1060,170))
    draw.text((445,490),'是',font=small,fill='black')
    arrow((155,365),(65,365));arrow((65,365),(65,785));arrow((65,785),(190,785))
    draw.text((75,320),'否',font=small,fill='black')
    draw.rectangle((190,735,1310,835),outline='black',width=3)
    draw.text((250,764),'任一校验失败：拒绝操作并提示原因，不输出单票明文',font=small,fill='black')
    img.save(output)


def revise(role):
    path = next((ROOT/'reports').glob('成员'+role+'*.docx'))
    original = subprocess.check_output(['git','show',f'{BASELINE}:reports/{path.name}'],cwd=ROOT)
    doc = Document(io.BytesIO(original))
    styles(doc)
    measure = json.loads((EVIDENCE/'full/experiment_results.json').read_text('utf-8'))
    config = ROLE_CONTENT[role]

    # Cover: the original school identity, table borders and evaluation form remain.
    for i in (12,13):
        p=doc.paragraphs[i]
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.left_indent=Pt(54 if i==12 else 153)
        ind=p._element.get_or_add_pPr().find(qn('w:ind'))
        if ind is not None:
            for key in ('firstLineChars','leftChars','hangingChars'):
                ind.attrib.pop(qn('w:'+key),None)
        for run in p.runs: font(run,22,False,'黑体','黑体')
    for row_id in (1,2,5):
        for p in doc.tables[0].rows[row_id].cells[1].paragraphs:
            for run in p.runs: font(run,15)

    # A small dedicated break paragraph avoids spilling an empty cover line onto
    # the evaluation sheet after its text fields are enlarged to template size.
    from docx.text.paragraph import Paragraph
    between=[]
    el=doc.tables[0]._tbl.getnext()
    while el is not None and el is not doc.tables[1]._tbl:
        between.append(el);el=el.getnext()
    for el in between:el.getparent().remove(el)
    el=OxmlElement('w:p');doc.tables[1]._tbl.addprevious(el)
    pagebreak=Paragraph(el,doc._body)
    pagebreak.paragraph_format.page_break_before=True
    pagebreak.paragraph_format.space_after=Pt(0)
    pagebreak.paragraph_format.space_before=Pt(0)
    pagebreak.paragraph_format.line_spacing=Pt(1)
    pagebreak.add_run().font.size=Pt(1)

    start=next(i for i,p in enumerate(doc.paragraphs) if p.text=='基于同态加密的隐私保护投票系统')
    body = doc.paragraphs[start:]
    for p in body:
        text=p.text
        is_picture=bool(p._element.xpath('.//w:drawing'))
        is_code=any(r.font.name=='Consolas' for r in p.runs)
        if re.match(r'^[一二三四五六] ', text):
            text=text[0]+'、'+text[2:]
            if text.startswith('四、'):text='四、部分核心代码'
            replace(p,text); name='Report H1'
        elif re.match(r'^\d\.\d ',text): name='Report H2'
        elif is_code:name='Report Code'
        elif text.startswith('图1'): name='Report Caption'
        elif text.startswith('基于同态加密的'):name='Report Title'
        elif any(r.bold for r in p.runs):name='Report H2'
        else:name='Report Body'
        p.style=name
        # Clear old direct formatting so the reviewed role styles actually apply.
        pp=p._element.get_or_add_pPr()
        for child in list(pp):
            if child.tag != qn('w:pStyle'): pp.remove(child)
        if name=='Report Title':p.paragraph_format.page_break_before=True
        if is_picture:
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            if is_picture:continue
            if name=='Report Code':font(run,10,False,'新宋体','Consolas')
            else:font(run,{'Report H1':14,'Report Title':16,'Report Caption':10.5}.get(name,12),name in ('Report H1','Report H2','Report Title'))
        if text.startswith('成员'):
            p.paragraph_format.space_after=Pt(6)

    def find(prefix):
        return next(p for p in doc.paragraphs if p.text.startswith(prefix))
    def set_text(prefix,text):replace(find(prefix),text)

    second=find('二、')
    h=insert_after(second,'2.1 本课题的主要功能','Report H2')
    h=insert_after(h,config['functions'])
    h=insert_after(h,'2.2 本课题的设计过程','Report H2')
    h=insert_after(h,config['design'])
    h=insert_after(h,'2.3 模块接口与分工','Report H2')
    insert_after(h,'表1 小组模块划分','Report Caption').paragraph_format.keep_with_next=True
    set_text('共同接口以JSON', '公开选票由A交给C登记，C把签名清单与密文交给B汇总，B向保管端提交完整汇总请求。各角色报告分别讨论编码与密钥、汇总与解密、界面与审计；实际姓名、学号和完成记录由成员核对填写。')

    # Every independently submitted report explains the full cryptographic path.
    if role=='A':
        p=find('选取不同的大素数')
        replace(p,'密钥生成取不同大素数p、q，令n=pq、g=n+1；私钥可用λ=lcm(p−1,q−1)与μ表示，其中μ是L(g^λ mod n²)的模n逆元，L(u)=(u−1)/n。加密为c=g^m r^n mod n²，r与n互素；解密为m=L(c^λ mod n²)μ mod n。phe实现以p、q和CRT加速，教学公式不等同于源码逐行实现。')
        insert_after(p,'同态关系为E(a;r₁)·E(b;r₂) mod n² = E((a+b) mod n;r₁r₂ mod n)。逐列相乘可累计各选项的0或1，最后只解密各列总票。参与人数上限10000远小于n，不会产生票数模n回绕。该关系由原语接口[2]及E(2)·E(3)解密为5的实测共同核对。')
    elif role=='B':
        insert_before(find('Paillier满足'),'上游先生成n=pq及g=n+1，把单选编码为独热向量；每项m以c=g^m r^n mod n²独立随机加密。私钥p、q由口令保管库保护。汇总端只读取n，保管端才加载p、q并调用raw_decrypt，这与原语接口的分工一致[2]。')
    else:
        p=find('清单payload包含')
        insert_before(p,'算法全流程：A生成Paillier公钥n=pq、g=n+1及私钥p、q，以c=g^m r^n mod n²加密独热向量；B按选项列相乘，利用E(a)·E(b) mod n² = E((a+b) mod n)，再由保管端加载私钥解密各列总票[2]。C的Ed25519签名用于发现文件变化，不能代替加密或证明独热合法性[3]。')

    # Replace prose arrows with an actual decision/exception flowchart.
    flow=find('3.4 模块处理流程')
    replace(flow,'3.4 模块程序流程图设计')
    nxt=flow._p.getnext()
    if nxt is not None and nxt.tag==qn('w:p'):
        from docx.text.paragraph import Paragraph
        q=Paragraph(nxt,flow._parent)
        if '→' in q.text:q._element.getparent().remove(q._element)
    work=ROOT/'work/report-assets';work.mkdir(parents=True,exist_ok=True)
    diagram=work/f'flow-{role}.png';flowchart(role,diagram)
    p=insert_after(flow,'','Report Caption');p.add_run().add_picture(str(diagram),width=Cm(16))
    p.paragraph_format.keep_with_next=True
    insert_after(p,'图1 本模块处理流程与失败出口','Report Caption')
    set_text('图1 '+{'A':'加密投票','B':'独立进程','C':'临时副本'}[role],
             '图2 '+{'A':'加密投票页的公开元数据','B':'独立进程返回总票数及实测耗时','C':'临时副本上的篡改和删除检测'}[role])
    set_text('模块界面与数据说明','3.6 模块界面与数据说明')
    set_text('公开数据与私有数据','3.7 公开数据与私有数据')
    set_text('性能数据与结果分析','5.3 性能数据与结果分析')
    set_text('准确性与可重复性','5.4 准确性与可重复性')
    set_text('额外功能的实际意义','5.5 本轮发现的问题与处理')
    issue=find('5.5 ')
    following=__import__('docx.text.paragraph',fromlist=['Paragraph']).Paragraph(issue._p.getnext(),issue._parent)
    replace(following,config['issue'])

    env=measure['environment']
    set_text('实测日期为','本轮实测日期为2026年10月27日（北京时间）。环境为'+env['os']+'，Python '+env['python']+'；phe 1.5.0、gmpy2 2.3.2、cryptography 50.0.2、PySide6-Essentials 6.12.0。核心实验使用2048位Paillier模数；9项工程回归使用1024位密钥加速边界检查，不作为安全强度示范。')
    set_text('自动核心实验共','自动核心实验29项全部符合预期，另有9项工程回归和原生Qt界面测试通过。原始记录位于evidence/2026-10-27/full/，包含JSON与CSV；源码SHA-256逐文件记入JSON。表2列出本模块相关结果。“局限”用例通过表示成功复现边界，不表示已防御。')
    set_text('对10、100、1000名投票者','对10、100、1000名投票者各做3次三选项实验。合成明文使用种子20261026加重复序号，密码学随机数不固定；每次均核对解密计数与预期。表3取本轮JSON中的算术平均值，单位毫秒。')
    set_text('计时范围仅包含','计时仅覆盖Paillier原语，不包括文件、签名、口令派生、进程及GUI；界面耗时则来自当前轮次的实际审计和解密端复算，二者口径不同。测试期间计算机有其他任务，三次原始记录保存在JSON中，均值不能当作部署保证。')
    set_text('基准6票的选择序列','基准6票序列[0,1,0,2,0,1]得到[3,2,1]；2、4、8选项各投一票均得到全1向量，零票返回全0。原生Qt测试记录在evidence/2026-10-27/ui/，验证6次控件投票、子进程计票和两种审计演示。图2采用本轮实际截图。')
    set_text('本项目的初始代码','本次使用OpenAI Codex辅助审查代码、修订报告、执行实验和拆分Git提交。初稿与本轮修改均有仓库记录，AI参与范围详见根目录AI使用声明.txt。A/B/C为模块责任标识，学生需据真实贡献补全个人信息和复核记录，并独立解释代码。')
    if role=='C':
        set_text('1．启动MiPiaoLab.exe','1．按README安装依赖并运行src/main.py；创建页填写标题、逗号分隔选项、人数及保管口令，选择目录创建新轮次。')
        set_text('实际修复：源码界面','界面复测：本轮在Windows原生Qt平台执行实际按钮点击，6票结果为[3,2,1]，随后篡改与删除演示均被检测。旧版线程等待处理在此次流程中继续有效；对应记录为ui/ui_smoke.json和截图。')

    # Update measured tables while retaining their existing geometry.
    for table in doc.tables[2:]:
        head=table.cell(0,0).text
        if head=='人数':
            for row, data in zip(table.rows[1:],measure['performance']):
                values=[str(data['voters'])]+[f'{data[k]:.2f}' for k in ('encrypt_ms','aggregate_ms','decrypt_ms')]
                for c,value in zip(row.cells,values):c.text=value
        if head=='实验项':
            for title,detail in config['extra']:
                row=table.add_row();row.cells[0].text=title;row.cells[1].text=detail
        for i,row in enumerate(table.rows):
            row.height=None
            if i==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
            for cell in row.cells:
                borders=cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
                if borders is None:
                    borders=OxmlElement('w:tcBorders');cell._tc.get_or_add_tcPr().append(borders)
                for edge in ('top','left','bottom','right'):
                    line=OxmlElement('w:'+edge);line.set(qn('w:val'),'single')
                    line.set(qn('w:sz'),'4');line.set(qn('w:color'),'D9D9D9');borders.append(line)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.2
                    if head in ('成员','人数'):
                        p.paragraph_format.keep_with_next=(i < len(table.rows)-1)
                    p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.space_before=Pt(2)
                    for run in p.runs:font(run,10.5,i==0)
        if head in ('实验项','人数'):
            from docx.text.paragraph import Paragraph
            el=OxmlElement('w:p');table._tbl.addprevious(el)
            cap=Paragraph(el,table._parent);cap.style='Report Caption'
            cap.add_run('表2 本模块测试结果' if head=='实验项' else '表3 加密 汇总与解密实测均值')
            cap.paragraph_format.keep_with_next=True

    # Correct report code line wrapping; the executable source remains authoritative.
    if role == 'A':
        block = [p for p in doc.paragraphs if p.style.name == 'Report Code']
        snippet = """    if type(choice) is not int or not 0 <= choice < len(config['options']):
        raise VoteError('必须且只能选择一个有效选项')
    if voter_id not in config['roster']:
        raise VoteError('投票者ID不在授权名单内')
    pub = public_key(config)
    encrypted = [
        str(pub.raw_encrypt(int(i == choice)))
        for i in range(len(config['options']))
    ]
    return {
        'version': 1, 'election_id': config['election_id'],
        'voter_id': voter_id, 'ballot_type': 'single_choice',
        'timestamp': now(), 'nonce': uuid.uuid4().hex,
        'ciphertexts': encrypted,
    }"""
        replace(block[0], snippet)
        for p in block[1:]:p._element.getparent().remove(p._element)
    for p in doc.paragraphs:
        if p.text.startswith('以下片段摘自'):
            p.paragraph_format.keep_with_next=True
        if p.text.startswith('['):
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=1.2
            for r in p.runs:font(r,10.5)
        if any(token in p.text for token in ('g^m','g^λ','c^λ','r^n')):
            chunks=re.split(r'([gcr]\^[mλn])',p.text)
            p.clear()
            for chunk in chunks:
                if re.fullmatch(r'[gcr]\^[mλn]',chunk):
                    p.add_run(chunk[0]);p.add_run(chunk[2]).font.superscript=True
                else:p.add_run(chunk)
        if p.style.name=='Report Code':
            if 'encrypted = [str(pub.raw_encrypt' in p.text:
                replace(p,"    encrypted = [\n        str(pub.raw_encrypt(int(i == choice)))\n        for i in range(len(config['options']))\n    ]")
            for r in p.runs:font(r,10,False,'新宋体','Consolas')
            p.paragraph_format.keep_with_next=True
    code=[p for p in doc.paragraphs if p.style.name=='Report Code']
    if code:code[-1].paragraph_format.keep_with_next=False
    footer=doc.sections[0].footer.paragraphs[0]
    footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.add_run(' 页')
    for r in footer.runs:font(r,9)
    doc.save(path)

    # Replace only the pre-existing UI screenshot part, preserving the school logo.
    photo=EVIDENCE/'ui'/({'A':'02_voting.png','B':'03_results.png','C':'04_audit.png'}[role])
    with ZipFile(io.BytesIO(original)) as z:
        media=[x for x in z.namelist() if x.startswith('word/media/') and x.lower().endswith('.png')]
        screenshot=max(media,key=lambda x:len(z.read(x)))
    with ZipFile(path) as z: parts={x:z.read(x) for x in z.namelist()}
    editable={'word/document.xml','word/styles.xml','word/footer1.xml',
              'word/_rels/document.xml.rels','[Content_Types].xml',screenshot}
    with ZipFile(io.BytesIO(original)) as z:
        for name in z.namelist():
            if name not in editable:parts[name]=z.read(name)
    parts[screenshot]=photo.read_bytes()
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    print(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--role',required=True,choices=['A','B','C'])
    args=parser.parse_args()
    revise(args.role)
