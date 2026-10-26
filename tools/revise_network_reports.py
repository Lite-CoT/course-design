"""Rebuild the three reports' bodies in their preserved template for the LAN implementation."""
import argparse
import io
import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from revise_reports import styles, font, flowchart, ROLE_CONTENT

ROOT=Path(__file__).resolve().parents[1]
EVIDENCE=ROOT/'evidence/2026-10-27-network'
BASELINE='ebd220f'
CONTENT={
 'A':{
  'name':'选票编码、客户端加密与密钥分离',
  'files':'crypto_service.py、network_setup.py、network_client.py',
  'function':'投票人在客户端输入IP和端口，导入组织者分配的独立凭证。客户端核验服务器证书指纹、配置摘要和绑定ID后，读取选项并在本端完成随机加密。初始化程序把服务器运行材料、保管端私钥和人员凭证分目录交付。',
  'design':'首先把原来集中在同一电脑上的密钥职责拆开。Paillier私钥仅写入trustee目录的口令保管库；server目录只接收公钥、登记签名密钥及TLS密钥。客户端独立计算选票，HTTPS请求体不包含choice字段，也不包含私钥或保管口令。随后增加证书指纹固定和配置摘要检查，避免把凭证交给错误服务器。',
  'detail':'凭证包含随机票码、绑定ID、轮次、配置摘要和TLS证书指纹。连接过程先建立TLS并比较证书SHA-256，再发送票码；取得公开配置后还要与凭证摘要比较。加密调用只依赖公钥n与用户选择，返回密文向量。独立随机数使相同选择产生不同密文，公开nonce只用于记录唯一性，不能作为Paillier内部随机数r。',
  'issue':'从单机改为联网时，最容易出现“服务器接收choice后再加密”的形式改造。此次把encrypt_ballot调用放在NetworkClient.vote中，接口仅接受ballot对象。测试截取客户端实际发送的对象，并比较服务器SQLite内的内容，确认没有明文选择字段；另外检查服务器文件中不存在Paillier私钥p、q及保管端结果签名私钥。',
  'flow':['输入IP、端口并导入本人凭证','证书与配置校验通过？','选择一个选项并构造独热向量','在本端独立随机加密各分量','HTTPS只发送密文与元数据','核对服务器签名回执及摘要'],
  'screenshot':'01_connected.png',
  'figure':'投票端输入IP、端口并验证凭证后取得选项',
  'code':"""started = perf_counter()
ballot = encrypt_ballot(
    self.election['config'], self.ticket['voter_id'], choice)
encrypted = perf_counter()
response = self.request(
    'POST', '/api/ballots', {'ballot': ballot})
received = perf_counter()
receipt = response['receipt']
payload = verify_object(
    receipt, self.election['config']['audit_public_key'])""",
  'code_source':'network_client.py中NetworkClient.vote，按版面折行；后续还核对ID、轮次和密文摘要',
  'questions':'为何不能把choice发送到服务器后再加密？为何公开nonce不能复用为r？证书指纹与配置摘要分别保护什么？凭证被转交后，服务器还能证明操作者身份吗？',
  'conclusion':'本模块将个人选择与Paillier私钥留在各自端点，建立了实际的客户端加密入口。服务端只得到随机密文；凭证分发和端点可信仍是前提。后续应完善凭证撤销与密文独热合法性证明。',
 },
 'B':{
  'name':'服务器同态汇总、离线解密与签名结果',
  'files':'network_protocol.py、network_trustee.py；与network_store.py的导出、发布接口协作',
  'function':'服务器审计已登记的密文集合，在关闭状态按选项列相乘，输出完整公开汇总包。独立保管端使用本地可信配置复核集合及乘积，再读取本地私钥解密各选项总票并签名。服务器核验签名后公布结果，所有客户端可验证。',
  'design':'计票分为服务器密文计算与保管端总票解密两个步骤。服务器进程不导入解密操作，也没有接收私钥的接口。保管端先验证关闭清单，再独立计算完整乘积，拒绝单票或部分汇总请求；使用批准记录把同一轮次绑定到一个确定集合，结果签名密钥与服务器登记签名密钥分离。',
  'detail':'聚合向量的初值为1，逐列执行C_j=∏c_ij mod n²，因此空集合解密为零。保管端对选票ID、随机标识、摘要、顺序和签名清单作一致性核验，并比较自己计算的密文向量与请求。结果绑定轮次、配置摘要、清单摘要和汇总包摘要；只有本地批准记录允许的完整集合才返回结果。重复处理同一个汇总返回已有签名结果。',
  'issue':'原来的独立解密子进程仍与收票端共享同一电脑上的私钥文件，无法证明网络服务器没有私钥。此次把解密入口改为离线decrypt命令，服务器只接收签名结果。测试还构造由服务器合法重签的另一个子集合，验证保管端在已批准一个汇总后拒绝再次解密不同集合；这仍不能替代首次集合的外部见证。',
  'flow':['接收关闭后的完整密文汇总包','可信配置与清单一致？','独立逐列复算并核对密文汇总','检查批准记录并加载本地私钥','解密总票、验证约束并签名','服务器发布，各客户端验签'],
  'screenshot':'04_published.png',
  'figure':'客户端验证保管端签名后显示公开总票',
  'code':"""n = int(election['config']['n'])
totals = [1] * len(election['config']['options'])
for ballot in ballots:
    totals = [
        (value * int(c)) % (n*n)
        for value, c in zip(totals, ballot['ciphertexts'])
    ]
aggregate = {
    'election_id': election['config']['election_id'],
    'manifest_sha256': digest(payload),
    'accepted': len(ballots),
    'ciphertexts': [str(value) for value in totals],
}""",
  'code_source':'network_protocol.py中aggregate_records，省略前置审计和计时，按版面折行',
  'questions':'服务器怎样在没有私钥时得到密文总票？为什么聚合初值为1？为什么关闭前不提供阶段总票？一次批准记录能防止什么，为什么仍需要外部回执或见证？',
  'conclusion':'本模块形成了服务器无私钥计算、保管端独立核验、客户端验证签名结果的完整路径。计算正确性通过已知票数及重算确认；保管者可信、首次集合完整性与单人结果隐私仍是边界。',
 },
 'C':{
  'name':'联网收票、事务审计与可视化交互',
  'files':'network_store.py、network_server.py、network_ui.py、network_cli.py、network_ui_smoke.py',
  'function':'服务器开放可配置的HTTPS端口；客户端输入IP、端口并导入凭证连接。接收端把高熵票码绑定ID，在SQLite事务中检查状态和重复记录，原子写入选票及签名清单。桌面界面显示验证后的回执、接收数量和公开结果，管理员具有关闭、导出与发布操作。',
  'design':'原单机JSON文件与QLockFile只适合一个写入者。联网后使用BEGIN IMMEDIATE串行化写事务，voter_id主键与nonce唯一约束共同限制重复登记。HTTPS处理器负责消息大小、JSON格式、凭证及角色检查；存储模块负责事务与完整性。Qt后台线程处理网络和加密，主线程更新控件。',
  'detail':'一次登记同时写入密文选票和新的签名清单，成功提交后才返回签名回执。管理员关闭与投票登记使用同一事务机制，避免关闭状态和接收记录交叉。客户端断线后可重新查询本人回执。服务器重启时先核验数据库中的清单与选票，发现篡改即拒绝正常服务。普通投票凭证不能关闭轮次或发布结果。',
  'issue':'真实并发会暴露“先检查再写入”的竞态，不能只依赖界面禁用按钮。本轮用8个TLS客户端同时使用同一ID提交，实测只接受1次、其余7次收到冲突拒绝。首次测试还发现SQLite连接的with块只提交事务、不会关闭连接，导致Windows清理测试目录失败；补充显式关闭后全部回归通过。',
  'flow':['HTTPS接收凭证与密文选票','身份、轮次与格式合法？','事务内检查状态及ID是否已投','原子写入选票和签名清单','关闭后服务器按列汇总密文','验证保管端签名并公布总票'],
  'screenshot':'03_admin.png',
  'figure':'管理员使用独立凭证连接同一服务器',
  'code':"""db = sqlite3.connect(
    self.root/'votes.sqlite3', timeout=15,
    isolation_level=None)
try:
    db.execute('BEGIN IMMEDIATE' if write else 'BEGIN')
    yield db
    db.commit()
except Exception:
    db.rollback()
    raise
finally:
    db.close()""",
  'code_source':'network_store.py中transaction上下文管理器，按版面折行；accept在此事务内完成检查与两项写入',
  'questions':'文件锁为何不足以处理多端投票？检查重复与登记为何必须在同一事务内？连接中断后怎样确认选票是否收到？证书指纹错误时是否已经发送票码？',
  'conclusion':'本模块把实验算法组织为可连接的局域网服务，并用事务、凭证权限和签名回执处理多端协作。当前已完成同机真实网络与多进程验证；跨物理电脑仍需核验网络互通和防火墙设置，不能以本机测试代替。',
 }
}


def revise(role):
    path=next((ROOT/'reports').glob(f'成员{role}*.docx'))
    original=subprocess.check_output(['git','show',f'{BASELINE}:reports/{path.name}'],cwd=ROOT)
    doc=Document(io.BytesIO(original))
    for name in ('Report H1','Report H2'):
        for el in doc.styles[name].element.xpath('./w:pPr/w:outlineLvl'):
            el.getparent().remove(el)
    styles(doc)
    start=next(p for p in doc.paragraphs if p.text=='基于同态加密的隐私保护投票系统')
    node=start._p
    while node is not None and node.tag!=qn('w:sectPr'):
        following=node.getnext();node.getparent().remove(node);node=following
    item=CONTENT[role]
    measurement=json.loads((EVIDENCE/'source/network-processes.json').read_text('utf-8'))
    old_perf=json.loads((ROOT/'evidence/2026-10-27/full/experiment_results.json').read_text('utf-8'))
    def p(text, style='Report Body', page=False):
        para=doc.add_paragraph(text,style)
        if page: para.paragraph_format.page_break_before=True
        return para
    def h(text):return p(text,'Report H2')
    def caption(text):
        para=p(text,'Report Caption');para.paragraph_format.keep_with_next=True;return para
    def table(headers, rows, widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,label in enumerate(headers):t.cell(0,i).text=label
        for row in rows:
            cells=t.add_row().cells
            for c,text in zip(cells,row):c.text=str(text)
        for i,row in enumerate(t.rows):
            if i==0:
                flag=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(flag)
            for j,cell in enumerate(row.cells):
                if widths:cell.width=Cm(widths[j])
                pr=cell._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'C9D4DF');border.append(el)
                pr.append(border)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF3');pr.append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing=1.2
                    para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:font(run,10.5,i==0)
        return t
    p('基于同态加密的隐私保护投票系统','Report Title',True)
    h(f'成员{role}分工报告  {item["name"]}')
    p('一、课题内容和要求','Report H1')
    h('1.1 课题内容')
    p('设计一个电子投票系统，利用同态加密技术，使服务器在不查看个人投票内容的情况下统计总票。本实现采用Paillier加法同态：投票端本地加密，服务器只接收密文并进行列乘积，独立保管端只解密完整汇总，服务器公布经过签名的总票。')
    p('操作场景为局域网服务器开放端口，各客户端填写IP和端口连接，也允许本机启动多个客户端模拟多端。选票包含ID、票型、时间戳、随机标识与密文向量；通过凭证绑定和签名记录检测重复、越权及数据修改。')
    h('1.2 本分工要求')
    p('本分工负责'+item['name']+'，需要说明数据输入、输出、失败条件及与其他模块的接口；通过可复跑的实验验证功能、耗时和安全边界。')
    p('本分工主要代码：'+item['files']+'。A/B/C表示模块责任，姓名学号与实际个人贡献由成员据实确认。')
    p('二、需求分析和总体设计','Report H1')
    h('2.1 本课题的主要功能');p(item['function'])
    h('2.2 本课题的设计过程');p(item['design'])
    h('2.3 系统部署与接口')
    p('组织者在保管端初始化三个目录：server部署到收票电脑，trustee留在独立保管电脑，organizer内的凭证按身份分发。初始化过程中产生私钥，但运行中的网络服务器从未获得Paillier私钥；服务器TLS私钥和登记签名私钥仅用于通信、审计，不具有选票解密能力。')
    caption('表1 三端职责和数据边界')
    table(['角色','输入与操作','可持有的秘密'],[
        ['客户端','IP、端口、凭证；本地加密后提交密文','本人票码、当前选择'],
        ['服务器','认证登记、关闭、密文汇总与结果发布','TLS和登记签名密钥、票码摘要'],
        ['保管端','离线核验完整汇总、解密并签署总票','Paillier私钥、结果签名私钥']], [2.4,8.0,5.2])
    p('网络接口均以/api为前缀，以JSON交换数据：election和status返回公开状态；ballots只接受密文选票；me返回绑定身份和本人回执；管理员使用close、export和results。服务器不提供接收明文choice、私钥或保管口令的接口。')
    p('三、相关功能模块详细设计','Report H1')
    h('3.1 同态加密原理与计票过程')
    p('Paillier密钥取不同素数p、q，令n=pq、g=n+1。明文m使用与n互素的独立随机数r加密：c=g^m r^n mod n²。私钥可用λ=lcm(p−1,q−1)与μ表示，解密为m=L(c^λ mod n²)μ mod n，其中L(u)=(u−1)/n；phe实现利用p、q及CRT加速[2]。')
    p('同态关系为E(a;r₁)E(b;r₂) mod n² = E((a+b) mod n;r₁r₂ mod n)。单选编码为独热向量，服务器对每个选项分别累乘密文，得到计数之和的加密形式。人数上限10000远小于2048位n，因此正常计数不会模n回绕。服务器计算列乘积无需解密。')
    h('3.2 本模块设计');p(item['detail'])
    h('3.3 验证与失败处理')
    p('票码使用系统随机数生成并在服务器仅保存SHA-256摘要。TLS证书指纹固定在凭证中，客户端先验指纹后发送票码。接收端检查字段、轮次、ID、密文维度与合法密文空间；关闭与登记在SQLite事务内处理。保管端结果绑定完整汇总摘要，服务器和客户端分别核验Ed25519签名[3]。')
    h('3.4 模块程序流程图设计')
    ROLE_CONTENT[role]['flow']=item['flow']
    assets=ROOT/'work/network-report-assets';assets.mkdir(parents=True,exist_ok=True)
    diagram=assets/f'{role}-flow.png';flowchart(role,diagram)
    picture=p('','Report Caption');picture.add_run().add_picture(str(diagram),width=Cm(16));picture.paragraph_format.keep_with_next=True
    p('图1 本模块处理流程及失败出口','Report Caption')
    h('3.5 模块界面与角色隔离')
    picture=p('','Report Caption');picture.add_run().add_picture(str(EVIDENCE/'ui'/item['screenshot']),width=Cm(16));picture.paragraph_format.keep_with_next=True
    p('图2 '+item['figure'],'Report Caption')
    p('图2来自本轮Windows原生Qt实际运行。投票人和管理员导入不同凭证，连接同一端口后得到各自权限。实际部署只向服务器复制server目录，保管端通过公开汇总文件与签名结果文件交接；同机测试使用独立目录和进程，不能声称已建立物理隔离。')
    p('四、部分核心代码','Report H1')
    h('4.1 核心实现片段');p('以下片段摘自'+item['code_source']+'。完整异常处理和调用上下文以源码为准。').paragraph_format.keep_with_next=True
    code=p(item['code'],'Report Code');code.paragraph_format.keep_together=True
    h('4.2 可重复实验步骤')
    for text in [
      '1．按照README安装依赖，在保管者电脑运行init生成新轮次。分发server目录、各人凭证；保留trustee目录和口令。',
      '2．服务器运行serve指定目录、监听IP和端口。客户端启动默认界面，填写服务器IPv4地址和端口，导入本人凭证并连接。本机多端使用127.0.0.1。',
      '3．使用不同ID依次或同时投票，记录纸面预期计数。重复ID、名单冒用和错误指纹应被拒绝，关闭前不能导出汇总。',
      '4．管理员关闭并导出密文汇总；保管端运行decrypt，输入本地口令，产生签名总票；管理员发布后，各端刷新核对结果。',
      '5．重新连接检查回执，重启服务器检查状态；再初始化不同人数和选项的新轮次。跨电脑时核对防火墙与网络互通，现场结果单独记录。']:
        p(text)
    p('五、课题完成过程中遇到的问题及解决方法','Report H1')
    h('5.1 原实现问题与本轮修正');p(item['issue'])
    p('单机阶段已有真实密码学运算，但没有网络收票服务器，因此原整体完成判断不充分。本次把联网架构作为主要实现，原单机入口仅保留为算法实验。实验记录和报告分别标注所属阶段，避免用旧单机测试证明多端可用。')
    h('5.2 实验环境与结果')
    p('本轮日期为2026年10月27日（北京时间），Windows 11 x64、Python 3.12.14，phe 1.5.0、gmpy2 2.3.2、cryptography 50.0.2、PySide6-Essentials 6.12.0。端到端验收采用2048位Paillier；18项联网边界测试使用1024位密钥加速，不作为安全强度示范。另有9项既有工程回归。')
    caption('表2 联网与交互验证结果')
    table(['项目','实际结果'],[
      ['TLS与权限、事务、完整性','18项联网测试＋9项回归全部通过'],
      ['6个独立客户端进程','实际连接同一TLS端口，6票结果[3,2,1]'],
      ['同一ID并发8次','只接受1次，其余7次拒绝'],
      ['凭证冒用、错误指纹','被拒绝；指纹失败时未发送票码'],
      ['错误口令、部分汇总、伪造结果','被拒绝，不发布错误总票'],
      ['两窗口实际Qt控件','6次投票，回执与公开结果验签通过'],
      ['服务重启','已关闭状态和已发布结果保持一致']], [5.0,10.6])
    p('网络与GUI测试均使用虚拟ID和合成选票，原始证据在evidence/2026-10-27-network/。独立进程验收在一台电脑上通过真实TLS完成，不代表跨物理电脑实测；教师私有数据和现场小改仍待实际验收。')
    h('5.3 加密、汇总和解密耗时')
    enc=[row['client_metrics']['encrypt_ms'] for row in measurement['receipts']]
    submit=[row['client_metrics']['submit_roundtrip_ms'] for row in measurement['receipts']]
    caption('表3 本轮六端并发实测（毫秒）')
    table(['测量范围','数值'],[
      ['每端一票的本地加密，六端算术均值',f'{sum(enc)/len(enc):.2f}'],
      ['提交密文至收到回执，六端均值',f'{sum(submit)/len(submit):.2f}'],
      ['服务器6票列乘积（不含审计）',f'{measurement["server_aggregation_ms"]:.2f}'],
      ['保管端3列总票解密',f'{measurement["trustee_decrypt_ms"]:.2f}']], [11.5,4.1])
    p('加密计时包含独热向量构造与三分量随机加密，提交往返包含TLS、服务器事务和响应；汇总与解密计时只覆盖各自运算。六端在同一机器并发，包含CPU竞争；不能把一次运行当作网络部署性能保证。逐端原始值、进程耗时与源码哈希均保存在network-processes.json。')
    caption('表4 历史原语性能对照（各三次均值，毫秒）')
    table(['人数','加密','汇总','解密'],[[str(row['voters'])]+[f'{row[k]:.2f}' for k in ('encrypt_ms','aggregate_ms','decrypt_ms')] for row in old_perf['performance']], [2.4,4.4,4.4,4.4])
    p('表4来自本次架构修正前保留的完整原语实验evidence/2026-10-27/full/，仅作规模对照。它不包含联网、事务和TLS开销，也不是重新执行的联网千人测试；原始运行环境和哈希在其JSON中保留。')
    h('5.4 失败、误报与安全局限')
    p('文件损坏、网络断开、错误凭证与攻击可能出现相似提示，程序只报告可验证的失败原因。数据库内容与签名不一致时拒绝继续；连接中断时通过回执确认是否已接收，而不是直接重新计票。首次打包复测发现Windows重定向输出使用GBK，已统一管道输出UTF-8，再通过实际EXE进程验证。')
    p('未实现逐票独热零知识证明，[2,0]与[0,0]等非法组合可能仍满足总票数约束。单人结果会暴露选择。一次批准只限制同轮重复解密不同集合；若恶意服务器重新签署删票集合，首次集合是否完整仍需独立回执或外部见证，不能只靠服务器自己的签名。')
    p('六、总结','Report H1')
    p(item['conclusion'])
    h('6.1 威胁模型与后续改进')
    p('主要模型是服务器按协议收票与汇总、但可能试图了解个人选择；在Paillier私钥留在可信保管端、客户端可信的条件下，服务器无法直接解密单票。该模型不覆盖客户端中毒、私钥泄露、组织者错误分发凭证或服务器与保管者串通。后续可增加独热证明、门限解密、独立透明日志、凭证撤销与抗拒绝服务措施。')
    h('6.2 分工验收与个人复核')
    p(item['questions']);p('成员应在自己的电脑复跑相关实验，补全姓名学号、实际贡献与本人修改记录，并现场解释代码、完成教师要求的小改。当前分工是模块划分，不代替个人完成证明。')
    h('6.3 AI使用声明')
    p('本项目使用OpenAI Codex辅助需求对照、代码与界面实现、联网架构修正、测试执行、报告修订及Git提交。测试数据来自实际运行的合成选票，AI参与范围详见根目录AI使用声明.txt；教师评分、签名和个人复核记录由实际参与者填写。')
    h('参考资料')
    for text in [
      '[1] 专业课程设计任务书，B230414，2026—2027学年第一学期。',
      '[2] python-paillier API. https://python-paillier.readthedocs.io/en/stable/phe.html',
      '[3] Cryptography Ed25519. https://cryptography.io/en/latest/hazmat/primitives/asymmetric/ed25519/',
      '[4] Python sqlite3. https://docs.python.org/3/library/sqlite3.html',
      '[5] Python ssl. https://docs.python.org/3/library/ssl.html']:
        para=p(text);para.alignment=WD_ALIGN_PARAGRAPH.LEFT;para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.line_spacing=1.2
        for run in para.runs:font(run,10.5)
    for para in doc.paragraphs:
        if para.style.name=='Report Code':
            for run in para.runs:font(run,10,False,'新宋体','Consolas')
        if any(token in para.text for token in ('g^m','r^n','c^λ')):
            chunks=re.split(r'([gcr]\^[mλn])',para.text);para.clear()
            for chunk in chunks:
                if re.fullmatch(r'[gcr]\^[mλn]',chunk):para.add_run(chunk[0]);para.add_run(chunk[2]).font.superscript=True
                else:para.add_run(chunk)
    doc.save(path)
    with ZipFile(path) as archive:parts={name:archive.read(name) for name in archive.namelist()}
    editable={'word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
    with ZipFile(io.BytesIO(original)) as archive:
        for name in archive.namelist():
            if name not in editable:parts[name]=archive.read(name)
    with ZipFile(path,'w',ZIP_DEFLATED) as archive:
        for name,data in parts.items():archive.writestr(name,data)
    print(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--role',choices=['A','B','C'],required=True)
    revise(parser.parse_args().role)
